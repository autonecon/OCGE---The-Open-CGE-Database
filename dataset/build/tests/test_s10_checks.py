"""Mutation tests for the S10 hard checks, the S10 screens and the config stops.

Every hard check is called on a good input (it must pass) and on a mutated input (it must
fail). A check that cannot fail on its mutation is not a check. The screens suite runs on a
small fake build; each screen must report its mutation. The config loaders must stop on a
missing, mis-keyed or incomplete file.

Run from dataset/build:  python3 -m unittest tests.test_s10_checks -v
"""
import csv
import io
import os
import shutil
import sys
import tempfile
import unittest
import zipfile
from collections import defaultdict
from types import SimpleNamespace

BUILD = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BUILD)

from s10_tariffs import checks as K  # noqa: E402
from s10_tariffs import config  # noqa: E402
from s10_tariffs.config import ConfigError  # noqa: E402
from s10_tariffs import preferences as prefs  # noqa: E402
from s10_tariffs.sources import Schedule  # noqa: E402
import screens.s10_suite as suite  # noqa: E402


def passes(result):
    return result[1] is True


class Nomen:
    """Identity schedule code: the tests use HS22 codes as schedule codes."""

    def code(self, nomen, hs, has=None):
        return hs


def sched_code(nomen, hs):
    return hs


def agreement(partner, family, year, rates, members):
    return {'partner': partner, 'family': family, 'year': year, 'nomen': 'H6', 'rates': rates,
            'rate_year': {k: year for k in rates}, 'members': set(members), 'excluded': {}}


def state_with(*agrs, imp='AAA'):
    cover = defaultdict(list)
    for a in agrs:
        for m in a['members']:
            cover[m].append(a)
    return {'agreements': {imp: list(agrs)}, 'cover': {imp: dict(cover)}, 'excl': {}, 'log': defaultdict(list)}


class HardChecks(unittest.TestCase):
    def test_join(self):
        good = {'baci': [], 'reporters': [], 'members': ['MYT'], 'numeric': ['234']}
        self.assertTrue(passes(K.join(good, {'MYT'})))
        for key, val in (('baci', [('999', 'XXX', 'XXX', 5.0)]), ('reporters', ['QQQ']), ('members', ['ZZZ']), ('numeric', ['111'])):
            bad = dict(good, **{key: val})
            self.assertFalse(passes(K.join(bad, {'MYT'})), key)

    def test_groups_and_unknown_groups(self):
        self.assertTrue(passes(K.groups([])))
        self.assertFalse(passes(K.groups(['A36 member list differs'])))
        self.assertTrue(passes(K.unknown_groups({})))
        self.assertFalse(passes(K.unknown_groups({'unknown_group': [('AAA', 'X99', 2021)]})))

    def test_above_mfn(self):
        self.assertTrue(passes(K.above_mfn([])))
        self.assertFalse(passes(K.above_mfn([('AAA', 'BBB', '010101', 5.0, 4.0)])))

    def test_best_applied(self):
        st = state_with(agreement('P1', 'f1', 2021, {'010101': 2.0}, {'BBB'}))
        tariffs = SimpleNamespace(state=st, sched_code=sched_code)
        good = [('AAA', 'BBB', '010101', 2.0, 5.0, 'preference'), ('DEU', 'FRA', '010101', 0.0, 5.0, 'cu_forced')]
        self.assertTrue(passes(K.best_applied(good, tariffs)))
        for row in (('AAA', 'BBB', '010101', 3.0, 5.0, 'preference'),     # above the family line
                    ('AAA', 'BBB', '010101', 6.0, 5.0, 'mfn'),            # above MFN
                    ('DEU', 'FRA', '010101', 1.0, 5.0, 'cu_forced')):     # forced zero not zero
            self.assertFalse(passes(K.best_applied([row], tariffs)), row)

    def test_family_spots(self):
        st = state_with(agreement('A36', 'eu', 2017, {}, {'DEU'}), agreement('A84', 'cefta', 2021, {}, {'DEU'}), imp='BIH')
        self.assertTrue(passes(K.family_spots(st, [('BIH', 'DEU', ('A36', 'A84'))])))
        st['cover']['BIH']['DEU'] = st['cover']['BIH']['DEU'][:1]
        self.assertFalse(passes(K.family_spots(st, [('BIH', 'DEU', ('A36', 'A84'))])))

    def test_khm_partner(self):
        st = state_with(agreement('356', 'acfta', 2021, {}, {'CHN'}), imp='KHM')
        self.assertTrue(passes(K.khm_partner(st)))
        self.assertFalse(passes(K.khm_partner(state_with(agreement('356', 'acfta', 2021, {}, {'IND'}), imp='KHM'))))

    def test_us_gsp(self):
        self.assertTrue(passes(K.us_gsp(0)))
        self.assertFalse(passes(K.us_gsp(1)))

    def test_uk_continuity(self):
        st = state_with(agreement('EUN', 'eu', 2021, {}, {'GBR', 'FRA'}), imp='MEX')
        self.assertTrue(passes(K.uk_continuity(st, {'MEX'}, {'EUN'})))
        self.assertFalse(passes(K.uk_continuity(st, set(), {'EUN'})))
        self.assertFalse(passes(K.uk_continuity(state_with(imp='DZA'), {'DZA'}, {'EUN'})))

    def test_schedule_values(self):
        s = Schedule('KOR', 'H6', 2022, 'test')
        s.set('120190', 4.5, 'korea_soy')
        expected = [('KOR', '120190', 4.5, 'korea_soy')]
        self.assertTrue(passes(K.schedule_values({'KOR': s}, Nomen(), expected)))
        s.set('120190', 487.0, 'ave')
        self.assertFalse(passes(K.schedule_values({'KOR': s}, Nomen(), expected)))
        self.assertFalse(passes(K.schedule_values({}, Nomen(), expected)))

    def test_iraq(self):
        self.assertTrue(passes(K.iraq(0)))
        self.assertFalse(passes(K.iraq(3)))

    def test_weight_years(self):
        good = {2021: [100.0, 99.95, 0.05, 0.0], 2022: [110.0, 110.0, 0.0, 0.0], 2023: [120.0, 120.0, 0.0, 0.0]}
        self.assertTrue(passes(K.weight_years(good)))
        for bad in ({k: v for k, v in good.items() if k != 2023},                      # a year missing
                    {**good, 2021: [100.0, 90.0, 5.0, 0.0]},                           # value lost
                    {**good, 2021: [100.0, 99.0, 1.0, 0.0]},                           # too much dropped
                    {**good, 2020: [1.0, 1.0, 0.0, 0.0]}):                             # an extra year
            self.assertFalse(passes(K.weight_years(bad)), bad)

    def test_weight_stream_splits_hs17(self):
        from s10_tariffs import weights as W
        tmp = tempfile.mkdtemp()
        p17, p22 = os.path.join(tmp, 'a.csv'), os.path.join(tmp, 'b.csv')
        with open(p17, 'w') as fh:
            fh.write('t,i,j,k,v,q\n2021,4,8,10101,9.0,1\n2021,4,8,999999,1.0,1\n2021,4,8,20202,4.0,1\n')
        with open(p22, 'w') as fh:
            fh.write('t,i,j,k,v,q\n2022,4,8,010101,3.0,1\n')
        corr = {'010101': ['010101'], '020202': ['020201', '020209']}
        wy = W.WeightYears([(2021, 'hs17', p17, None), (2022, 'hs22', p22, None)], corr, {'020201': 3.0, '020209': 1.0})
        rows = list(wy.flows())
        got = defaultdict(float)
        for _i, _j, hs, v, _q, _y, w in rows:
            got[hs] += v * w
        self.assertAlmostEqual(got['010101'], 9.0 / 2 + 3.0 / 2)
        self.assertAlmostEqual(got['020201'], 3.0 / 2)
        self.assertAlmostEqual(got['020209'], 1.0 / 2)
        self.assertEqual(wy.stats[2021], [14.0, 13.0, 1.0, 0.0])
        self.assertTrue(passes(K.weight_years(wy.stats, years=(2021, 2022), max_dropped=0.1)))

    def test_ncl_schedule_and_exemption(self):
        from s10_tariffs import rates as R, sources
        tmp = tempfile.mkdtemp()
        path = os.path.join(tmp, 'ncl.csv')
        with open(path, 'w') as fh:
            fh.write('code8,hs6,dd_text,dd_pct,tgc_text,other_taxes\n01012100,010121,EX,0.0,3%,\n'
                     '01012910,010129,10%,10.0,3%,\n01012990,010129,20%,20.0,3%,\n')
        p22 = os.path.join(tmp, 'ncl22.csv')
        with open(p22, 'w') as fh:
            fh.write('hs6,dd_2022_pct,dd_2023_pct,lines_2022,lines\n')
        sch = sources.ncl_schedule(path, p22)
        self.assertEqual(sch.rate, {'010121': 0.0, '010129': 15.0})
        with open(p22, 'a') as fh:
            fh.write('010121,5.0,0.0,2,x\n')
        self.assertEqual(sources.ncl_schedule(path, p22).rate['010121'], 5.0)      # 2022 DD restored
        for bad in ('010129,5.0,10.0,2,x\n', '010199,5.0,10.0,2,x\n'):          # 2023 average or HS6 does not match
            with open(p22, 'w') as fh:
                fh.write('hs6,dd_2022_pct,dd_2023_pct,lines_2022,lines\n' + bad)
            with self.assertRaises(ValueError):
                sources.ncl_schedule(path, p22)
        with open(p22, 'w') as fh:
            fh.write('hs6,dd_2022_pct,dd_2023_pct,lines_2022,lines\n')
        state = {'agreements': {}, 'cover': {}, 'excl': {}}
        unions = SimpleNamespace(kind=lambda i, e, h: None)
        t = R.Tariffs({'NCL': sch}, state, Nomen(), unions, {}, set())
        self.assertEqual(t.line('NCL', 'AUS', '010129'), (15.0, 15.0, 'reported'))
        self.assertEqual(t.line('NCL', 'FRA', '010129'), (0.0, 15.0, 'ncl_eu_oct'))
        self.assertEqual(t.line('NCL', 'PYF', '010129')[0], 0.0)
        with open(path, 'a') as fh:
            fh.write('01013000,010130,RM5/kg,,3%,\n')
        with self.assertRaises(ValueError):
            sources.ncl_schedule(path, p22)

    def test_tunisia_turkey_lists(self):
        from s10_tariffs import rates as R
        tmp = tempfile.mkdtemp()
        path = os.path.join(tmp, 'l.csv')
        with open(path, 'w') as fh:
            fh.write('list,national_code,hs6_2002,source_file\nI,25151220105,251512,x\nII,61091000001,610910,x\n'
                     'III,57011010001,570110,x\nII,57011010002,570110,x\nII,44011000001,440110,x\nI,61091000002,610910,x\n')
        lists = R.load_tun_lists(path)
        self.assertEqual(lists, {'251512': (1, 0, 0), '610910': (1, 1, 0), '570110': (0, 1, 1), '440110': (0, 1, 0)})
        tun = Schedule('TUN', 'H6', 2016, 'test')
        for hs in ('251512', '610910', '570110', '440110'):
            tun.set(hs, 30.0, 'reported')
        unions = SimpleNamespace(kind=lambda i, e, h: None)
        t = R.Tariffs({'TUN': tun}, {'agreements': {}, 'cover': {}, 'excl': {}}, Nomen(), unions, {}, set(),
                      turkey_rule=True, tun_lists=lists)
        self.assertEqual(t.line('TUN', 'TUR', '251512')[0], 0.0)      # List I: abolished by 2012
        self.assertAlmostEqual(t.line('TUN', 'TUR', '610910')[0], 9.0)   # ch.61, one List I and one List II line: (0 + 18) / 2
        self.assertEqual(t.line('TUN', 'TUR', '440110')[0], 0.0)      # List II outside the Finance Law chapters
        self.assertEqual(t.line('TUN', 'TUR', '570110')[0], 15.0)     # List III line at MFN, List II line outside the chapters at 0
        with open(path, 'w') as fh:
            fh.write('list,national_code,hs6_2002,source_file\nI,25151220105,251512,x\n')
        with self.assertRaises(ConfigError):
            R.load_tun_lists(path)

    def test_national_preferences(self):
        from s10_tariffs import sources
        tmp = tempfile.mkdtemp()
        path = os.path.join(tmp, 'm.csv')
        with open(path, 'w') as fh:
            fh.write('agreement,partner,nomen,national_code,hs6,rate_text,rate_pct,column,source_file\n'
                     'ACFTA,156,H5,0101210000,010121,0,0.0,c,f\nACFTA,156,H5,0101290010,010129,5,5.0,c,f\n'
                     'ACFTA,156,H5,0101290090,010129,15,15.0,c,f\nACFTA,156,H5,0101300000,010130,N.O.,,c,f\n')
        agreements = {'ACFTA': ('156', frozenset({'CHN'}))}
        recs, stats = sources.national_preferences(path, agreements)
        self.assertEqual(stats, {'ACFTA': (4, 3)})
        self.assertEqual(len(recs), 1)
        self.assertEqual(recs[0]['rates'], {'010121': 0.0, '010129': 10.0})
        self.assertEqual((recs[0]['iso'], recs[0]['partner'], recs[0]['year'], recs[0]['source']), ('MYS', '156', 2022, 'national'))
        self.assertEqual(recs[0]['members'], frozenset({'CHN'}))
        with self.assertRaises(ConfigError):                           # partner code differs from the table
            sources.national_preferences(path, {'ACFTA': ('410', frozenset({'CHN'}))})
        with self.assertRaises(ConfigError):                           # extract agreement without a table row
            sources.national_preferences(path, {})
        with self.assertRaises(ConfigError):                           # table row without extracted lines
            sources.national_preferences(path, dict(agreements, RCEP=('B39', frozenset({'CHN'}))))
        # members from the table, not from the group list: RCEP code B39 with Malaysia's 2022 parties
        merged = {('MYS', 'B39', 'H5'): {'iso': 'MYS', 'partner': 'B39', 'nomen': 'H5', 'year': 2022, 'rates': {'010121': 0.0},
                                         'rate_year': {}, 'excluded': {}, 'members': frozenset({'AUS', 'KOR'}),
                                         'agreement': 'RCEP'}}
        state = prefs.build_state(merged, {'B39': ('RCEP ASEAN', ['BRN', 'KHM'])}, set(), set(), [], {'B39': 'rcep'}, {}, None)
        agr = state['agreements']['MYS'][0]
        self.assertEqual((agr['members'], agr['family'], agr['name']), ({'AUS', 'KOR'}, 'rcep', 'RCEP'))

    def test_national_agreements_table(self):
        class CM:
            def iso(self, code):
                return code if len(code) == 3 and code.isalpha() else ''
        tmp = tempfile.mkdtemp()
        path = os.path.join(tmp, 'a.csv')
        head = 'agreement,partner_code,members_2022,entry_into_force_for_mys,source,decided_by,decided_at\n'
        with open(path, 'w') as fh:
            fh.write(head + 'RCEP,B39,AUS KOR,2022-03-18,s,David,2026-10-23\n')
        self.assertEqual(prefs.load_national_agreements(CM(), path), {'RCEP': ('B39', frozenset({'AUS', 'KOR'}))})
        for bad in ('RCEP,B39,AUS KOR,2023-01-02,s,David,2026-10-23',   # not in force in 2022
                    'RCEP,B39,AUS 410,2022-03-18,s,David,2026-10-23',   # member that does not map
                    'RCEP,B39,AUS,2022-03-18,s,,2026-10-23',            # no decided_by
                    'RCEP,,AUS,2022-03-18,s,David,2026-10-23'):         # no partner code
            with open(path, 'w') as fh:
                fh.write(head + bad + '\n')
            with self.assertRaises(ConfigError, msg=bad):
                prefs.load_national_agreements(CM(), path)

    def test_weight_conversion(self):
        good = {'elc': (96.0, 100.0), 'mac': (104.0, 100.0), 'crp': (50.0, 50.0)}
        self.assertTrue(passes(K.weight_conversion(good)))
        for bad in ({'elc': (159.0, 100.0), 'mac': (41.0, 100.0), 'crp': (50.0, 50.0)},   # n:n split into new headings
                    {'elc': (100.0, 100.0), 'mac': (100.0, 100.0), 'crp': (44.0, 50.0)},  # one large sector 12% off
                    {'elc': (90.0, 100.0), 'dropped': (10.0, 0.0), 'crp': (50.0, 50.0)}):  # value lost
            self.assertFalse(passes(K.weight_conversion(bad, max_misallocation=0.02)), bad)

    def test_weight_convert_partition(self):
        from s10_tariffs import weights as W
        wy = W.WeightYears([None], {'847989': ['847989', '852499']}, {'847989': 30.0, '852499': 10.0})
        self.assertEqual(wy.convert('847989'), [('847989', 0.75), ('852499', 0.25)])
        self.assertEqual(wy.convert('271000'), None)
        tmp = tempfile.mkdtemp()
        path = os.path.join(tmp, 'c.xlsx')
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'HS2022-HS2017 Conversions'
        for r in (('From HS2022', 'To HS2017'), ('852411', '847989'), ('847989', '847989'), ('852411', '870899')):
            ws.append(r)
        wb.save(path)
        with self.assertRaises(ValueError):                             # an HS2022 code with two sources
            W.hs17_to_hs22(path)

    def test_parent_territory(self):
        terr = {'SMR': None}
        good = [('SMR', '010101', 6.0, 6.0, 0.0), ('SMR', '720851', None, None, None)]
        self.assertTrue(passes(K.parent_territory(good, terr)))
        for bad in ([('SMR', '010101', 4.0, 6.0, 4.0)],                 # MFN is not the parent line
                    [('SMR', '010101', 6.0, 6.0, 8.0)],                 # applied above MFN
                    [('SMR', '720851', 3.0, None, 3.0)],                # rate outside the union
                    [('SMR', '720851', None, None, None)]):             # nothing priced
            self.assertFalse(passes(K.parent_territory(bad, terr)), bad)

    def test_parent_schedule_line(self):
        from s10_tariffs import rates as R
        eun = Schedule('EUN', 'H6', 2022, 'test')
        eun.set('010101', 6.0, 'reported')
        eun.set('720851', 2.0, 'reported')
        eun.set('300490', 0.0, 'reported')
        eu_pref = agreement('A01', 'ch', 2022, {'010101': 1.0}, ['CHE'])
        state = {'agreements': {}, 'cover': {'AUT': {'CHE': [eu_pref]}}, 'excl': {}}
        unions = SimpleNamespace(kind=lambda i, e, h: None)
        parents = {'SMR': ('EUN', 'EUN', 1, ('72',)), 'AND': ('EUN', 'EUN', 25, ())}
        t = R.Tariffs({'EUN': eun}, state, Nomen(), unions, {}, set(), parents=parents)
        self.assertEqual(t.line('SMR', 'USA', '010101'), (6.0, 6.0, 'parent_schedule'))
        self.assertEqual(t.line('SMR', 'CHE', '010101'), (1.0, 6.0, 'preference'))
        self.assertEqual(t.line('SMR', 'USA', '720851')[0], None)        # outside the union
        self.assertEqual(t.line('AND', 'USA', '010101')[0], None)        # chapters 1-24
        self.assertEqual(t.line('AND', 'USA', '300490'), (0.0, 0.0, 'parent_schedule'))
        self.assertEqual(t.parent_line('SMR', '720851'), None)
        self.assertTrue(t.related('SMR', 'CHE'))

    def test_d_by_basis(self):
        from s10_tariffs.support import d_by_basis
        b = {'PHR': 100.0, 'PHR@XDC': 1000.0, 'PHR_AREA@XDC': 600.0, 'PHR_AN@XDC': 300.0, 'PHR_R@XDC': 100.0}
        land, cap, lab = d_by_basis(b)
        self.assertAlmostEqual(land, 60.0)
        self.assertAlmostEqual(cap, 30.0)
        self.assertAlmostEqual(lab, 10.0)
        self.assertIsNone(d_by_basis(dict(b, **{'PHR_R@XDC': 0.0})))        # split does not add up
        self.assertIsNone(d_by_basis({'PHR': 5.0}))                           # no split at all
        self.assertEqual(d_by_basis({}), (0.0, 0.0, 0.0))

    def test_basis_split(self):
        from s10_tariffs.support import basis_split
        b = {'PCAAN': 100.0, 'PCAAN@XDC': 1000.0, 'PC_AREA@XDC': 750.0, 'PC_AN@XDC': 250.0}
        self.assertEqual(basis_split(b, 'PCAAN', ('PC_AREA', 'PC_AN')), {'PC_AREA': 75.0, 'PC_AN': 25.0})
        self.assertIsNone(basis_split(dict(b, **{'PC_AN@XDC': 0.0}), 'PCAAN', ('PC_AREA', 'PC_AN')))   # does not add up
        self.assertIsNone(basis_split({'PCAAN': 5.0}, 'PCAAN', ('PC_AREA', 'PC_AN')))                 # no split
        self.assertEqual(basis_split({}, 'PCAAN', ('PC_AREA', 'PC_AN')), {'PC_AREA': 0.0, 'PC_AN': 0.0})

    def pse_run(self, eu_mapping, aan_split=None):
        from s10_tariffs import support as SP
        flags = []
        S = SimpleNamespace(flag=lambda *a, **k: flags.append(a[0]), adjust=lambda *a, **k: None)
        C = {k: config.CHOICES[k][0] for k in ('pse_eu', 'pse_mapping', 'pse_aan_split', 'pse_non_oecd', 'pse_eu_mapping')}
        C['pse_eu_mapping'] = eu_mapping
        if aan_split:
            C['pse_aan_split'] = aan_split
        blk = {'PO': 1.0, 'PIV': 2.0, 'PIF': 3.0, 'PIS': 0.0, 'PC': 20.0, 'PCAAN': 16.0, 'PCRI': 4.0, 'PHR': 0.0, 'PHNR': 50.0,
               'PCAAN@XDC': 16.0, 'PC_AREA@XDC': 12.0, 'PC_AN@XDC': 4.0, 'PCRI@XDC': 4.0, 'PC_R@XDC': 3.0, 'PC_INC@XDC': 1.0}
        mon = {'EU': dict(blk), 'USA': dict(blk)}
        fao = {('DEU', '2051'): 3.0, ('FRA', '2051'): 1.0, ('USA', '2041'): 1.0, ('USA', '2044'): 1.0,
               ('DEU', '2041'): 1.0, ('DEU', '2044'): 1.0, ('FRA', '2041'): 1.0, ('FRA', '2044'): 1.0}
        regs = ['DEU', 'FRA', 'USA']
        pre = {k: {r: 100.0 for r in regs} for k in ('crp_out', 'crp_inter', 'crp_va')}
        rows, booked, _, bad = SP.pse(S, C, mon, fao, regs, pre)
        out = defaultdict(float)
        for r in rows:
            out[(r[0], r[1])] += float(r[4])
        return out, booked, bad, flags

    def test_pse_eu_decoupled_support_on_land(self):
        out, booked, bad, flags = self.pse_run('ch13d_e_phnr_and_income_based_c_d_to_land_one_rate')
        self.assertEqual(bad, [])
        self.assertAlmostEqual(out[('DEU', 'factors')], 0.0)                     # ch.13D: no E on all factors in the EU
        self.assertAlmostEqual(out[('DEU', 'land')], 0.75 * (50.0 + 12.0 + 1.0))   # E + C area + C income, DEU share 3/4
        self.assertAlmostEqual(out[('DEU', 'labour')], 0.75 * 3.0)                # C receipts stay on labour
        self.assertAlmostEqual(out[('DEU', 'capital')], 0.75 * (3.0 + 4.0))       # PIF + C animal numbers (MONEVA basis)
        self.assertAlmostEqual(out[('USA', 'factors')], 50.0)                     # non-EU: ch.13D standard formula, E on all factors
        self.assertAlmostEqual(out[('USA', 'land')], 12.0)
        self.assertAlmostEqual(out[('USA', 'labour')], 4.0)
        for r in ('DEU', 'FRA', 'USA'):                                          # A2-E booked in full either way
            self.assertAlmostEqual(booked[r], (76.0 if r == 'USA' else 76.0 * (0.75 if r == 'DEU' else 0.25)))
        self.assertIn('PHNR-EU-LAND', flags)
        self.assertIn('PCAAN-BY-BASIS', flags)
        alt, *_ = self.pse_run('standard_formula_as_non_eu')                     # mutation: the standard formula in the EU
        self.assertAlmostEqual(alt[('DEU', 'factors')], 0.75 * 50.0)
        self.assertAlmostEqual(alt[('DEU', 'land')], 0.75 * 12.0)
        fao_split, *_ = self.pse_run('standard_formula_as_non_eu', config.CHOICES['pse_aan_split'][1])
        self.assertAlmostEqual(fao_split[('USA', 'land')], 8.0)                  # mutation: FAOSTAT 50/50 output share, not the basis

    def test_panama(self):
        tpa = {'010101': 2.0}
        self.assertTrue(passes(K.panama([('010101', 2.0, 10.0)], tpa)))
        self.assertFalse(passes(K.panama([('010101', 5.0, 10.0)], tpa)))
        self.assertFalse(passes(K.panama([], tpa)))   # nothing priced is a failure

    def test_ave_1000(self):
        lines = [('EGY', '220300', 1200.0, 5.0)]
        self.assertTrue(passes(K.ave_1000(lines, {('EGY', '220300')})))
        self.assertFalse(passes(K.ave_1000(lines, set())))

    def test_customs_union(self):
        cells = {('DEU', 'FRA', 'crp'): 0.0, ('DEU', 'USA', 'crp'): 3.0}
        self.assertTrue(passes(K.customs_union(cells, {('DEU', 'FRA', 'crp')})))
        self.assertFalse(passes(K.customs_union({**cells, ('DEU', 'FRA', 'crp'): 1.0}, {('DEU', 'FRA', 'crp')})))

    def test_no_addon_in_tm(self):
        base = {('USA', 'CHN', 'eeq'): 2.0}
        self.assertTrue(passes(K.no_addon_in_tm(dict(base), base)))
        self.assertFalse(passes(K.no_addon_in_tm({('USA', 'CHN', 'eeq'): 27.0}, base)))

    def test_weights(self):
        self.assertTrue(passes(K.weights(1000.0, 1000.0)))
        self.assertFalse(passes(K.weights(1000.0, 990.0)))

    def test_every_cell_rated(self):
        rules = {('AAA', 'BBB', 'crp'): 'bilateral', ('AAA', 'CCC', 'crp'): 'missing'}
        self.assertTrue(passes(K.every_cell_rated(rules, {'AAA'})))
        self.assertFalse(passes(K.every_cell_rated({k: 'missing' for k in rules}, {'AAA'})))

    def test_wedge_files(self):
        # B9: the written files against the S07+S08 inputs and the BACI accumulator, not against themselves
        from s10_tariffs.util import gz_write
        tmp = tempfile.mkdtemp()
        k, k2 = ('AAA', 'BBB', 'crp'), ('AAA', 'CCC', 'crp')

        def write(rows):
            gz_write(tmp + '/tm_s10.csv.gz', b'importer,exporter,sector,tm,rate_pct,fill,price_basis\n',
                     [f'{a},{b},{c},{tm},{tm * 100},bilateral,cif\n'.encode() for (a, b, c), (tm, _v, _m, _d) in rows.items()])
            gz_write(tmp + '/vmsb_s10.csv.gz', b'importer,exporter,sector,vcif_musd,vmsb_musd,vcif_basis,vmsb_basis\n',
                     [f'{a},{b},{c},{v},{m},cif,basic\n'.encode() for (a, b, c), (_t, v, m, _d) in rows.items()])
            gz_write(tmp + '/duty_bilateral_s10.csv.gz', b'importer,exporter,sector,duty_musd,vcif_musd,price_basis\n',
                     [f'{a},{b},{c},{d},{v},x\n'.encode() for (a, b, c), (_t, v, _m, d) in rows.items()])
            return K.read_back_cells(tmp)
        gz_write(tmp + '/vx.csv.gz', b'importer,exporter,sector,vfob_musd\n', [b'AAA,BBB,crp,90\n', b'AAA,CCC,crp,45\n'])
        gz_write(tmp + '/vt.csv.gz', b'importer,exporter,sector,vtwr_musd\n', [b'AAA,BBB,crp,10\n', b'AAA,CCC,crp,5\n'])
        vcif = K.vcif_independent(tmp + '/vx.csv.gz', tmp + '/vt.csv.gz', {'crp'})
        self.assertEqual(vcif[k], 100.0)
        good = {k: (0.05, 100.0, 105.0, 5.0), k2: (0.1, 50.0, 55.0, 5.0)}
        bil = {k: 5.0}
        self.assertTrue(passes(K.wedge_files(write(good), vcif, bil)))
        self.assertFalse(passes(K.wedge_files(write({**good, k: (0.06, 100.0, 106.0, 6.0)}), vcif, bil)))     # rate off the accumulator
        self.assertFalse(passes(K.wedge_files(write({**good, k: (0.05, 99.0, 103.95, 4.95)}), vcif, bil)))    # VCIF off S07+S08
        self.assertFalse(passes(K.wedge_files(write({k: good[k]}), vcif, bil)))                               # a VCIF cell not written
        self.assertFalse(passes(K.wedge_files(write({**good, k2: (0.1, 50.0, 55.0, 6.0)}), vcif, bil)))       # identity: duty
        self.assertFalse(passes(K.wedge_files(write({**good, k2: (-0.1, 50.0, 45.0, -5.0)}), vcif, bil)))     # identity: VMSB < VCIF

    def test_energy_tax_rows(self):
        """A1 (GTAP 12 ch.12 §12.3.3, §12.3.10): S10 writes each source region's class revenue with its base and the
        specific rate per GJ, and a proxy row for every other region; S10-Henergy on mutated rows."""
        from s10_tariffs import energy_tax as E
        classes = {'ROAD': ('p', 'rdt', 'dpabr'), 'INDUSTRY': ('pet',), 'RESCOM': ('p',), 'AGRIFISH': (), 'ELEC': ('ely',), 'OFFROAD': ()}
        rev = {'AAA': {('ROAD', 'pet'): [12.0, -2.0, 1000.0, 'TJ'], ('ALL', 'ely'): [1.0, 0.0, 0.0, 'electricity_s11']},
               'BBB': {('ROAD', 'pet'): [6.0, 0.0, 400.0, 'TJ'], ('RESCOM', 'gas'): [0.0, 0.0, 10.0, 'TJ']}}
        iso = {'AAA': 'AAA', 'BBB': 'BBB'}
        pools = {'AAA': 'P1', 'BBB': 'P1', 'CCC': 'P1', 'DDD': 'P2'}
        C = {'energy_tax_rows': config.CHOICES['energy_tax_rows'][0], 'energy_tax_fill': 'trade_weighted_s12_pool_average_else_world',
             'energy_tax_year': 'neer_2021'}
        regions = ['AAA', 'BBB', 'CCC', 'DDD']
        rows, summ = E.build(regions, iso, pools, {'AAA': 3.0, 'BBB': 1.0}, rev, classes, C)
        R = {(r[0], r[1], r[2]): dict(zip(E.FILE_COLUMNS, r)) for r in rows}
        road = R[('AAA', 'ROAD', 'pet')]
        self.assertAlmostEqual(float(road['specific_usd_per_gj']), 10.0 * 1000 / 1000.0)        # net USD m per TJ x 1000 = USD per GJ
        self.assertAlmostEqual(float(road['neer_subsidy_usd_per_gj']), -2.0)                    # B15: the NEERS subsidy part
        self.assertEqual(road['columns'], 'p;rdt')                                              # purchases abroad carry no energy tax
        self.assertEqual(R[('AAA', 'ALL', 'ely')]['specific_usd_per_gj'], '')                   # S13 divides by the S11 electricity use
        self.assertNotIn(('BBB', 'RESCOM', 'gas'), R)                                           # a zero revenue is not written
        self.assertEqual(R[('CCC', '*', '*')]['proxies'], 'AAA:3;BBB:1')                        # the pool's source regions, trade weights
        self.assertEqual(R[('DDD', '*', '*')]['proxies'], 'AAA:3;BBB:1')                        # no pool member: world (this choice)
        self.assertEqual(summ['filled'], ['CCC', 'DDD'])
        self.assertTrue(passes(K.energy_tax(rows, rev, iso, regions)))
        # GTAP 12 ch.12 fill (default): only pool members on the same side of energy trade; with none, zero
        Cg = dict(C, energy_tax_fill='same_energy_trade_direction_in_pool_else_zero')
        rg, sg = E.build(regions, iso, pools, {'AAA': 3.0, 'BBB': 1.0}, rev, classes, Cg, energy_net={'AAA': -5.0, 'BBB': 4.0, 'CCC': 9.0})
        Rg = {(r[0], r[1], r[2]): dict(zip(E.FILE_COLUMNS, r)) for r in rg}
        self.assertEqual(Rg[('CCC', '*', '*')]['proxies'], 'BBB:1')                             # the exporter BBB, not the importer AAA
        self.assertEqual(Rg[('DDD', '*', '*')]['proxies'], '')                                  # no pool member with data: zero
        self.assertIn('DDD', sg['zero'])
        self.assertTrue(passes(K.energy_tax(rg, rev, iso, regions)))
        self.assertEqual(config.CHOICES['energy_tax_fill'][0], 'same_energy_trade_direction_in_pool_else_zero')
        # mutations S10-Henergy must catch
        col = E.FILE_COLUMNS.index
        i = next(n for n, r in enumerate(rows) if r[:3] == ['AAA', 'ROAD', 'pet'])
        for k, v in (('tax_musd', '13'), ('specific_usd_per_gj', '12'), ('base', '900'), ('neer_subsidy_usd_per_gj', '0'), ('tax_musd', 'nan')):
            bad = [list(r) for r in rows]
            bad[i][col(k)] = v
            self.assertFalse(passes(K.energy_tax(bad, rev, iso, regions)), k)
        self.assertFalse(passes(K.energy_tax([r for n, r in enumerate(rows) if n != i], rev, iso, regions)))     # a revenue not written
        bad = [list(r) for r in rows]
        bad[-1][col('proxies')] = 'CCC:1'                                                        # a proxy that has no data
        self.assertFalse(passes(K.energy_tax(bad, rev, iso, regions)))
        self.assertFalse(passes(K.energy_tax(rows + [['AAA', '*', '*'] + [''] * 12], rev, iso, regions)))      # source and fill rows
        bad = [list(r) for r in rows]
        bad[next(n for n, r in enumerate(rows) if r[1] == 'ALL')][col('specific_usd_per_gj')] = '1'           # a per-GJ rate S10 cannot know
        self.assertFalse(passes(K.energy_tax(bad, rev, iso, regions)))
        rows0, _s = E.build(regions, iso, pools, {}, rev, classes, dict(C, energy_tax_rows='one_column_rate'))
        self.assertEqual(rows0, [])                                                              # the switch off: no energy rates
        self.assertNotIn('energy_tax_cap', config.CHOICES)                                       # the TLS cap is not an S10 rule any more
        self.assertIn('ch.12', E.Q_ENERGY)
        self.assertNotIn('ch.15', E.Q_ENERGY)

    def test_energy_tax_sources(self):
        from s10_tariffs import energy_tax as E
        neer = {(2023, 'AAA', 'ROAD', 'DIES'): {'TAXBTJ': 1000.0, 'TOTTAX': 10.0, 'TOTSUBS': -2.0},
                (2023, 'AAA', 'ELEC', 'COAL'): {'TAXBTJ': 500.0, 'FUETAX': 1.0, 'CARBTAX': 1.0, 'ELETAX': 4.0, 'SUBSID': 0.0},
                (2023, 'AAA', '_T', 'DIES'): {'TAXBTJ': 9e9, 'TOTTAX': 99.0},             # an aggregate is not read
                (2021, 'BBB', 'ROAD', 'DIES'): {'TAXBTJ': 1000.0, 'TOTTAX': 10.0}}
        rev = E.revenues(neer, {2021: 1.2, 2023: 1.1}, 'neer_2023_latest', 'net_of_neer_subsidies')
        self.assertEqual(set(rev), {'AAA'})                                       # BBB has no 2023 value
        self.assertAlmostEqual(rev['AAA'][('ROAD', 'pet')][0], 10.0 * 1000 * 1.1 / 1000)
        self.assertAlmostEqual(rev['AAA'][('ROAD', 'pet')][1], -2.0 * 1000 * 1.1 / 1000)
        self.assertAlmostEqual(rev['AAA'][('ALL', 'ely')][0], 4.0 * 500 * 1.1 / 1000)
        self.assertAlmostEqual(rev['AAA'][('ELEC', 'coa')][0], 2.0 * 500 * 1.1 / 1000)
        self.assertEqual(rev['AAA'][('ROAD', 'pet')][2:], [1000.0, 'TJ'])                 # the NEERS tax base: rate = tax / base
        self.assertEqual(rev['AAA'][('ELEC', 'coa')][2:], [500.0, 'TJ'])
        self.assertEqual(rev['AAA'][('ALL', 'ely')][3], 'electricity_s11')                # per GJ of electricity: S13 with S11 use
        self.assertEqual(E.revenues(neer, {2021: 1.2, 2023: 1.1}, 'neer_mean_2021_2023', 'taxes_only'), {})   # AAA lacks 2021, BBB 2023
        self.assertEqual(E.revenues(neer, {2021: 1.2, 2023: 1.1}, 'neer_2023_latest', 'taxes_only')['AAA'][('ROAD', 'pet')][1], 0.0)
        classes = E.load_classes()
        self.assertEqual(set(classes), set(E.NEER_SECTORS))
        self.assertIn('p', classes['ROAD'])
        self.assertNotIn('chm', classes['ROAD'])

    def test_neer_year_rule(self):
        from s10_tariffs import energy_tax as E
        self.assertEqual(config.CHOICES['energy_tax_year'][0], 'neer_2021')
        self.assertTrue(passes(K.neer_year(E.years_of('neer_2021'), '2021', 2022)))
        self.assertFalse(passes(K.neer_year(E.years_of('neer_2023_latest'), '2021', 2022)))     # after the base year
        self.assertFalse(passes(K.neer_year(E.years_of('neer_mean_2021_2023'), '2021', 2022)))
        self.assertFalse(passes(K.neer_year((2021,), 'NOT RECORDED', 2022)))                    # no source_years row
        yrs = {r['source_folder']: r for r in csv.DictReader(open(os.path.join(BUILD, 'config', 'source_years.csv'))) if r['base_year'] == '2022'}
        self.assertEqual(yrs['oecd_neer']['year_used'], '2021')
        for f in ('eia_us_motor_fuel_taxes', 'eia_foks', 'eurostat_energy_prices_households', 'wto_tpr'):
            self.assertIn(f, yrs)

    def test_usa_road_fuel_source(self):
        """Fix-round A1: the USA is a source region from EIA road-fuel excise (DF_NEERS has no US values)."""
        from s10_tariffs import energy_tax as E
        d = os.path.join(os.path.dirname(BUILD), 'data')
        paths = (d + '/eia_us_motor_fuel_taxes/pmm_fueltaxes.xlsx', d + '/eia_us_motor_fuel_taxes/MER_T03.07C.csv',
                 d + '/eia_foks/pet_cons_821use_dcu_nus_a.htm')
        rev, det = E.us_motor_fuel(*paths)
        self.assertAlmostEqual(det['federal_usd_per_gal']['gasoline'], 0.184, places=9)      # 26 USC 4081 + LUST
        self.assertAlmostEqual(det['federal_usd_per_gal']['diesel'], 0.244, places=9)
        self.assertAlmostEqual(det['rate_usd_per_gal']['gasoline'], 0.184 + (0.2605686274509804 + 0.2665) / 2, places=9)
        self.assertAlmostEqual(det['onhighway_share'], 40790502 / 45567273, places=12)
        self.assertEqual(det['onhighway_share_year'], 2020)
        self.assertAlmostEqual(det['gallons']['gasoline'], 8421.19 * 1000 * 42 * 365, places=0)
        self.assertEqual(set(rev), {('ROAD', 'pet')})
        self.assertAlmostEqual(rev[('ROAD', 'pet')][0], sum(det['tax_musd'].values()))
        self.assertEqual(rev[('ROAD', 'pet')][3], 'million_bbl')                             # the base is the taxed barrels
        self.assertAlmostEqual(rev[('ROAD', 'pet')][2], sum(det['gallons'].values()) / 42.0 / 1e6)
        self.assertTrue(70000 < rev[('ROAD', 'pet')][0] < 90000)
        # mutation: a PMM sheet whose columns moved stops the stage
        import openpyxl
        tmp = tempfile.mkdtemp()
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = 'January 2022'
            ws.append(['Federal and state motor fuels taxes'])
            ws.append([None, 'Excise', 'LUST Fee', None, None, 'Total'])               # Total one column off
            ws.append(['Federal', 0.183, 0.001, None, None, 0.184])
            ws.append([None, 'State tax'])
            ws.append(['Average state tax', 0.26])
            wb.create_sheet('July 2022 (revised)')
            p = os.path.join(tmp, 'bad.xlsx')
            wb.save(p)
            with self.assertRaises(config.ConfigError):
                E.read_pmm_motor_fuel(p)
        finally:
            shutil.rmtree(tmp)
        with self.assertRaises(config.ConfigError):
            E.read_mer_annual(paths[1], {'gasoline': 'MGACPUS'}, 1900)                   # no annual row for that year

    def test_usa_is_a_source_not_a_fill(self):
        from s10_tariffs import energy_tax as E
        classes = {'ROAD': ('p', 'rdt'), 'INDUSTRY': (), 'RESCOM': ('p',), 'AGRIFISH': (), 'ELEC': (), 'OFFROAD': ()}
        C = {'energy_tax_rows': config.CHOICES['energy_tax_rows'][0], 'energy_tax_fill': 'trade_weighted_s12_pool_average_else_world',
             'energy_tax_year': 'neer_2021'}
        rev = {'CAN': {('ROAD', 'pet'): [40.0, 0.0, 100.0, 'TJ'], ('RESCOM', 'pet'): [10.0, 0.0, 50.0, 'TJ']},
               'USA': {('ROAD', 'pet'): [20.0, 0.0, 2.0, 'million_bbl']}}
        rows, summ = E.build(['USA', 'CAN'], {'USA': 'USA', 'CAN': 'CAN'}, {'USA': 'NAMR', 'CAN': 'NAMR'}, {'USA': 1.0, 'CAN': 1.0}, rev,
                             classes, C, source_of={'USA': 'EIA test'})
        R = {(r[0], r[1], r[2]): dict(zip(E.FILE_COLUMNS, r)) for r in rows}
        self.assertNotIn('USA', summ['filled'])
        self.assertEqual(R[('USA', 'ROAD', 'pet')]['source'], 'EIA test')
        self.assertEqual(R[('USA', 'ROAD', 'pet')]['base_unit'], 'million_bbl')              # S13 converts barrels with the S11 factor
        self.assertNotIn(('USA', 'RESCOM', 'pet'), R)                                         # no RESCOM tax sourced
        # mutation: without the US source the USA is filled from its pool (CAN)
        rows2, summ2 = E.build(['USA', 'CAN'], {'USA': 'USA', 'CAN': 'CAN'}, {'USA': 'NAMR', 'CAN': 'NAMR'}, {'USA': 1.0, 'CAN': 1.0},
                               {'CAN': rev['CAN']}, classes, C)
        self.assertIn('USA', summ2['filled'])
        self.assertEqual({(r[0], r[1]): r[13] for r in rows2}[('USA', '*')], 'CAN:1')

    def test_korea_trq_measure_and_egypt_sublines(self):
        from s10_tariffs import overrides as ov
        plan = ov.korea_trq_rates()
        self.assertEqual(plan['100630']['regime'], 'in_quota')                    # decided: MAFRA plan fill 0.765
        baci = ov.korea_trq_rates('baci_2022_hs6_tonnes_over_quota')
        self.assertEqual(baci['100630']['regime'], 'out_of_quota')                # BACI 2022 tonnes over the quota: 1.18
        self.assertAlmostEqual(baci['100630']['rate'], 513.0)
        self.assertIn('baci_2022_hs6_tonnes_over_quota', config.CHOICES['trq_fill_measure'])
        self.assertEqual({k: v[0] for k, v in ov.EGY_MIXED_SUBLINES.items()}, {'330210': 5.0, '210690': 10.0})
        s = Schedule('EGY', 'H5', 2019, 'test')
        s.set('330210', 5.0, 'egypt_mixed_subline')
        expected = [('EGY', '330210', 5.0, 'egypt_mixed_subline')]
        self.assertTrue(passes(K.schedule_values({'EGY': s}, Nomen(), expected)))
        s.set('330210', 1003.3333, 'reported')                                   # the TRAINS HS6 average back on the line
        self.assertFalse(passes(K.schedule_values({'EGY': s}, Nomen(), expected)))

    def test_identical(self):
        self.assertTrue(passes(K.identical(['a', 'b'], ['a', 'b'])))
        self.assertFalse(passes(K.identical(['a', 'b'], ['a', 'c'])))
        self.assertFalse(passes(K.identical(['a', 'b'], ['a'])))

    def test_pse_total_preview(self):
        self.assertTrue(passes(K.pse_total([])))
        self.assertFalse(passes(K.pse_total(['USA booked 1 A2-E 2'])))
        self.assertTrue(passes(K.preview({'PO': 87.16}, {'PO': 87.1608}, 0.05)))
        self.assertFalse(passes(K.preview({'PO': 88.0}, {'PO': 87.1608}, 0.05)))
        self.assertFalse(passes(K.preview({}, {'PO': 87.1608}, 0.05)))

    def test_xs_rate_and_vxsb(self):
        row = ['KOR', 'crp', '2022', '1', '10', '2', '0.5', 'yes', 'own_hs', '', 'rice', 'fob']
        self.assertTrue(passes(K.xs_rate([row])))
        self.assertFalse(passes(K.xs_rate([row[:6] + ['1.5'] + row[7:]])))
        self.assertTrue(passes(K.vxsb(12.0, 12.0)))
        self.assertFalse(passes(K.vxsb(12.0, 11.0)))
        self.assertTrue(passes(K.vxsb(12.0, 12.0, 3.0, 3.0)))
        self.assertFalse(passes(K.vxsb(12.0, 12.0, 0.0, 3.0)))          # a booked export tax missing from VXSB

    def test_vfob_lines_export_tax(self):
        from s10_tariffs import support
        rows = [('crp', 'ARG', 'BRA', 100.0), ('crp', 'ARG', 'CHN', 100.0), ('fbt', 'ARG', 'CHN', 50.0)]
        lines, wedge, tax = support.vfob_lines(rows, {('ARG', 'fbt'): 50.0}, {('ARG', 'crp'): 200.0, ('ARG', 'fbt'): 50.0},
                                               {('ARG', 'crp'): 0.1})
        got = {tuple(x.decode().split(',')[:3]): float(x.decode().split(',')[3]) for x in lines}
        self.assertAlmostEqual(got[('crp', 'ARG', 'BRA')], 90.0)                 # VXSB = VFOB (1 - t)
        self.assertAlmostEqual(got[('fbt', 'ARG', 'CHN')], 100.0)                # VXSB = VFOB / (1 - s)
        self.assertAlmostEqual(tax, 20.0)
        self.assertAlmostEqual(wedge, 50.0)
        self.assertEqual(lines[0].decode().count(','), 7)                        # export_tax_rate column

    def test_export_tax_allocation_and_check(self):
        from s10_tariffs import export_tax as X
        products = {'ARG': [('*', ('1201',), 30.0, 'TPR p.1'), ('*', ('12',), 3.0, 'TPR p.2'), ('*', ('2709',), None, 'TPR p.3')],
                    'RUS': [('*', ('2709',), None, 'TPR p.4')]}
        exp_hs = {('ARG', '120190'): 1000.0, ('ARG', '120600'): 1000.0, ('ARG', '270900'): 1000.0, ('ARG', '271600'): 500.0}
        hs_sector = {'120190': 'crp', '120600': 'crp', '270900': 'oil', '271600': 'ely'}
        vfob = {('ARG', 'crp'): 2.0, ('ARG', 'oil'): 1.0}
        totals = {'ARG': 0.5, 'CIV': 1.0}
        rates, rows, rep = X.allocate(totals, products, exp_hs, hs_sector, vfob, {'ARG', 'RUS', 'CIV'})
        # weights: soya 30 (longest prefix), other crp 3, crude at the mean cited ad valorem rate 16.5; ely has no VFOB
        w_crp, w_oil = (30.0 + 3.0) / 100.0, 16.5 / 100.0
        self.assertAlmostEqual(rates[('ARG', 'crp')], 0.5 * w_crp / (w_crp + w_oil) / 2.0)
        self.assertAlmostEqual(rates[('ARG', 'oil')], 0.5 * w_oil / (w_crp + w_oil) / 1.0)
        self.assertNotIn(('ARG', 'ely'), rates)
        self.assertTrue(any('RUS' in x for x in rep['not_booked']))              # list, no 2022 revenue
        self.assertTrue(any('CIV' in x and 'NEEDS-SOURCE' in x for x in rep['not_booked']))
        self.assertTrue(passes(K.export_tax(rows, totals, products)))
        bad = [list(r) for r in rows]
        bad[0][3] = str(float(bad[0][3]) * 2)
        self.assertFalse(passes(K.export_tax(bad, totals, products)))            # duty no longer adds up to GRS
        bad = [list(r) for r in rows]
        bad[0][5] = '1.5'
        self.assertFalse(passes(K.export_tax(bad, totals, products)))            # rate above VFOB
        self.assertFalse(passes(K.export_tax(rows, totals, {})))                 # no cited product list

    def test_export_tax_products_table(self):
        from s10_tariffs import export_tax as X
        prods = X.load_products()
        self.assertTrue({'ARG', 'KAZ', 'IDN', 'MYS'} <= set(prods))
        self.assertTrue(all(src for rows in prods.values() for *_x, src in rows))
        tmp = tempfile.mkdtemp()
        path = os.path.join(tmp, 'p.csv')
        head = 'exporter,sector,hs_prefixes,rate_pct,rate_kind,product,source,page,decided_by,decided_at\n'
        for body in ('ARG,*,1201,33,ad_valorem,soy,TPR,,w3,2026-10-24\n', 'ARG,*,1201,33,guess,soy,TPR,1,w3,2026-10-24\n',
                     'ARG,*,,33,ad_valorem,soy,TPR,1,w3,2026-10-24\n'):
            with open(path, 'w') as fh:
                fh.write(head + body)
            with self.assertRaises(ConfigError):
                X.load_products(path)

    def test_export_tax_rates_path_and_check(self):
        # QA HIGH-2: RUS, IND and CHN booked from statutory 2022 rates (export_tax_rates_2022.csv)
        from s10_tariffs import export_tax as X
        tab = {'RUS': [('*', ('2709',), 0.08, 'crude'), ('*', ('271121',), 0.30, 'gas'), ('*', ('271111',), 0.0, 'LNG'),
                       ('*', ('2716',), 0.5, 'ely')]}
        exp_hs = {('RUS', '270900'): 1000.0, ('RUS', '271121'): 600.0, ('RUS', '271111'): 400.0, ('RUS', '271600'): 50.0,
                  ('RUS', '120190'): 10.0}
        hs_sector = {'270900': 'oil', '271121': 'gas', '271111': 'gas', '271600': 'ely', '120190': 'crp'}
        vfob = {('RUS', 'oil'): 1.2, ('RUS', 'gas'): 2.0, ('RUS', 'crp'): 0.01}
        rates, rows, rep = X.rates_booked(tab, exp_hs, hs_sector, vfob, {'RUS'}, {'RUS': 0.5})
        self.assertAlmostEqual(rates[('RUS', 'oil')], 0.08)
        self.assertAlmostEqual(rates[('RUS', 'gas')], 0.30 * 600.0 / 1000.0)          # LNG lines dilute the gas rate
        self.assertNotIn(('RUS', 'ely'), rates)                                          # no S07 VFOB to carry it
        self.assertNotIn(('RUS', 'crp'), rates)                                          # untaxed
        self.assertAlmostEqual(rep['revenue'][0][1], 0.08 * 1.2 + 0.18 * 2.0)
        self.assertEqual(rep['revenue'][0][2], 0.5)
        grs_rows = [['ARG', 'crp', '1.0', '1.0', '10.0', '0.1', '1', 'OECD GRS T_5124 2022', 'TPR', 'fob']]
        totals, products = {'ARG': 1.0}, {'ARG': [('*', ('12',), 3.0, 'TPR')]}
        self.assertTrue(passes(K.export_tax(grs_rows + rows, totals, products, tab, X.RATES_SOURCE)))
        bad = [list(r) for r in rows]
        bad[0][3] = str(float(bad[0][3]) * 1.5)
        self.assertFalse(passes(K.export_tax(bad, totals, products, tab, X.RATES_SOURCE)))   # duty is not rate x VFOB
        self.assertFalse(passes(K.export_tax(rows, totals, products, {}, X.RATES_SOURCE)))   # no rates-table row
        both = rows + [['RUS', 'crp', '1.0', '1.0', '10.0', '0.1', '1', 'OECD GRS T_5124 2022', 'TPR', 'fob']]
        self.assertFalse(passes(K.export_tax(both, {'RUS': 1.0}, {'RUS': []}, tab, X.RATES_SOURCE)))  # both paths
        bad = [list(r) for r in rows]
        bad[0][5] = '1.2'
        self.assertFalse(passes(K.export_tax(bad, totals, products, tab, X.RATES_SOURCE)))   # rate above VFOB
        # allocate leaves the rates-path exporters out
        _r, _rows, rep2 = X.allocate({'RUS': 5.0}, {'RUS': [('*', ('2709',), None, 'TPR')]}, exp_hs, hs_sector, vfob, {'RUS'},
                                     skip={'RUS'})
        self.assertEqual(_rows, [])
        self.assertEqual(rep2['not_booked'], [])

    def test_export_tax_rates_table(self):
        from s10_tariffs import export_tax as X
        tab = X.load_rates()
        self.assertTrue({'RUS', 'IND', 'CHN'} <= set(tab))
        got = {(iso, p): r for iso, rows in tab.items() for _s, pre, r, _src in rows for p in pre}
        self.assertAlmostEqual(got[('RUS', '271121')], 0.30)
        self.assertEqual(got[('RUS', '271111')], 0.0)
        self.assertTrue(0.05 < got[('RUS', '270900')] < 0.10)
        tmp = tempfile.mkdtemp()
        path = os.path.join(tmp, 'r.csv')
        head = ('exporter,sector,hs_prefixes,rate_pct,rate_kind,product,derivation,source,source_url,retrieved,decided_by,'
                'decided_at\n')
        good = 'RUS,*,271121,30,ad_valorem,gas,stat,Res 754,https://x,2026-10-24,wB,2026-10-24\n'
        with open(path, 'w') as fh:
            fh.write(head + good)
        self.assertAlmostEqual(X.load_rates(path)['RUS'][0][2], 0.30)
        for body in ('RUS,*,271121,30,ad_valorem,gas,stat,,https://x,2026-10-24,wB,2026-10-24\n',      # no source
                     'RUS,*,271121,30,guess,gas,stat,Res 754,https://x,2026-10-24,wB,2026-10-24\n',     # bad kind
                     'RUS,*,271121,100,ad_valorem,gas,stat,Res 754,https://x,2026-10-24,wB,2026-10-24\n',  # rate 100%
                     'RUS,*,271121,5,cited_zero,gas,stat,Res 754,https://x,2026-10-24,wB,2026-10-24\n',    # zero kind, rate 5
                     'RUS,*,,30,ad_valorem,gas,stat,Res 754,https://x,2026-10-24,wB,2026-10-24\n',         # no HS prefix
                     good + good):                                                                         # duplicate prefix
            with open(path, 'w') as fh:
                fh.write(head + body)
            with self.assertRaises(ConfigError):
                X.load_rates(path)

    def test_export_tax_rates_match_derivation(self):
        # config/s10/export_tax_rates_2022.csv carries exactly the rates of derived_aves_2022.csv
        from s10_tariffs import export_tax as X
        der_path = os.path.join(os.path.dirname(BUILD), 'data', 'export_duty_schedules_2022', 'derived_aves_2022.csv')
        with open(der_path, newline='') as fh:
            der = {(r['exporter'], r['hs_prefix']): float(r['ave_pct']) / 100.0 for r in csv.DictReader(fh)}
        cfg = {(iso, p): r for iso, rows in X.load_rates().items() for _s, pre, r, _src in rows for p in pre}
        self.assertEqual(set(cfg), set(der))
        for k, v in der.items():
            self.assertAlmostEqual(cfg[k], v, places=9, msg=str(k))

    def test_gfs_export_taxes_reader(self):
        from s10_tariffs import export_tax as X
        tmp = tempfile.mkdtemp()
        gfs, fx = os.path.join(tmp, 'g.csv'), os.path.join(tmp, 'fx.csv')
        with open(gfs, 'w') as fh:
            fh.write('COUNTRY,SECTOR,INDICATOR,TIME_PERIOD,OBS_VALUE\nRUS,S13,G1152_T,2022,2730231000000\n'
                     'RUS,S1311B,G1152_T,2022,1\nRUS,S13,G1151_T,2022,5\nCHN,S13,G1152_T,2022,0\nIND,S13,G1152_T,2022,\n')
        with open(fx, 'w') as fh:
            fh.write('series,iso3,year,value\nPA.NUS.FCRF,RUS,2022,68.4849416666666\nPA.NUS.FCRF,CHN,2022,6.7\n'
                     'PA.NUS.FCRF,IND,2022,78.6\n')
        got = X.read_gfs_export_taxes(gfs, fx)
        self.assertAlmostEqual(got['RUS'], 2730231.0 / 68.4849416666666, places=3)
        self.assertEqual(got['CHN'], 0.0)
        self.assertNotIn('IND', got)

    def test_listed(self):
        self.assertTrue(passes(K.listed(['a.csv'], {'a.csv'}, 'a tag')))
        self.assertFalse(passes(K.listed(['a.csv', 'b.csv'], {'a.csv'}, 'a tag')))

    def test_bind(self):
        self.assertTrue(passes(K.bind({'sums': ('abc', 'abc'), 'named': (True, True)})))
        self.assertFalse(passes(K.bind({'sums': ('abd', 'abc')})))
        self.assertFalse(passes(K.bind({'named': (False, True)})))

    def test_forbidden_inputs(self):
        self.assertTrue(passes(K.forbidden_inputs(['/d/trains_internal/x.tar'], 'value\n')))
        self.assertFalse(passes(K.forbidden_inputs(['/d/trains_internal/v3.8/x.csv'], '')))
        self.assertFalse(passes(K.forbidden_inputs([], 'old_plus_mfn' + '_change')))

    def test_vfob_pretax(self):
        self.assertTrue(passes(K.vfob_pretax(0)))
        self.assertFalse(passes(K.vfob_pretax(2)))

    def test_sources_present(self):
        good = {'TUN': 5000, 'TUN_raised': 10, 'LBY': 4000, 'VEN': 4000, 'TKM_scheduled': False}
        self.assertTrue(passes(K.sources_present(good)))
        for mut in ({'LBY': 0}, {'TUN': 500}, {'TUN_raised': 0}, {'TKM_scheduled': True}):
            self.assertFalse(passes(K.sources_present(dict(good, **mut))), mut)

    def test_spots(self):
        cells = {('MEX', 'USA', 'wat'): 0.0}
        mfn = {('MEX', 'USA', 'wat'): 10.0}
        self.assertTrue(passes(K.spots(cells, mfn, [('MEX', 'USA')], ('wat',))))
        self.assertFalse(passes(K.spots({('MEX', 'USA', 'wat'): 10.0}, mfn, [('MEX', 'USA')], ('wat',))))

    def test_pref_exceptions(self):
        rows = [{'importer': 'AGO', 'cause': 'none', 'source_checked': 'WITS', 'in_hpref_list': 'yes',
                 'trains_preference_data': 'no'}]
        self.assertTrue(passes(K.pref_exceptions(rows, {'MEX'})))
        self.assertFalse(passes(K.pref_exceptions(rows, {'MEX', 'AGO'})))
        self.assertFalse(passes(K.pref_exceptions([dict(rows[0], cause='')], {'MEX'})))
        self.assertFalse(passes(K.pref_exceptions(rows, set())))

    def test_tax_source_negatives_usa_far(self):
        self.assertTrue(passes(K.tax_source([])))
        self.assertFalse(passes(K.tax_source(['USA pit'])))
        rows = [('USA', 'PSE output', -5.0)]
        self.assertTrue(passes(K.negatives(0, rows, {('USA', 'PSE output')})))
        self.assertFalse(passes(K.negatives(0, rows, set())))
        self.assertFalse(passes(K.negatives(1, [], set())))
        self.assertTrue(passes(K.usa_far([])))
        self.assertFalse(passes(K.usa_far([29])))

    def test_screens_check(self):
        self.assertTrue(passes(K.screens({'material_unexplained': 0, 'detail': 'none'})))
        self.assertFalse(passes(K.screens({'material_unexplained': 2, 'detail': 'x'})))
        self.assertFalse(passes(K.screens({'error': 'KeyError'})))
        self.assertFalse(passes(K.screens(None)))

    def test_secret_and_private(self):
        import s10_tariffs_taxes as stage
        from lib.private_files import private_paths
        with tempfile.TemporaryDirectory() as d:
            clean = os.path.join(d, 'a.txt')
            with open(clean, 'w') as fh:
                fh.write('nothing here\n')
            self.assertEqual(stage.secret_hits([d]), [])
            with open(os.path.join(d, 'b.txt'), 'w') as fh:
                fh.write('password = ' + 'x' * 16 + '\n')
            self.assertEqual(stage.secret_hits([d]), ['b.txt'])
        self.assertEqual(private_paths(['dataset/data/baci/x.csv']), [])
        self.assertEqual(private_paths(['home/.secrets/wits.txt']), ['home/.secrets/wits.txt'])


class ConfigStops(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        with open(config.CHOICES_PATH, newline='', encoding='utf-8') as fh:
            self.rows = list(csv.DictReader(fh))
            self.cols = list(self.rows[0].keys())

    def write(self, rows, cols=None):
        path = os.path.join(self.tmp, 'c.csv')
        with open(path, 'w', newline='', encoding='utf-8') as fh:
            w = csv.DictWriter(fh, fieldnames=cols or self.cols, extrasaction='ignore')
            w.writeheader()
            w.writerows(rows)
        return path

    def test_good_file_loads(self):
        self.assertEqual(len(config.load_choices(self.write(self.rows))), len(config.CHOICES))

    def test_stops(self):
        rows = self.rows
        cases = {
            'missing file': os.path.join(self.tmp, 'absent.csv'),
            'missing row': self.write(rows[1:]),
            'unknown choice': self.write(rows + [dict(rows[0], choice='made_up_choice')]),
            'duplicate row': self.write(rows + [rows[0]]),
            'value not allowed': self.write([dict(rows[0], value='something_else')] + rows[1:]),
            'blank decided_by': self.write([dict(rows[0], decided_by='')] + rows[1:]),
            'bad gtap_status': self.write([dict(rows[0], gtap_status='maybe')] + rows[1:]),
            'missing column': self.write(rows, [c for c in self.cols if c != 'decided_at']),
        }
        for name, path in cases.items():
            with self.assertRaises(ConfigError, msg=name):
                config.load_choices(path)

    def test_every_choice_is_carried(self):
        """Every choice is named where the code carries it (C['name'] or a comment on the rule) or in
        DECISIONS_CARRIED, and DECISIONS_CARRIED names only real choices."""
        import s10_tariffs_taxes as stage
        src = ''
        for root, names in ((os.path.join(BUILD, 's10_tariffs'), None), (os.path.join(BUILD, 'screens'), ('s10_suite.py',))):
            for fn in os.listdir(root):
                if fn.endswith('.py') and fn != 'config.py' and (names is None or fn in names):
                    with open(os.path.join(root, fn), encoding='utf-8') as fh:
                        src += fh.read()
        with open(os.path.join(BUILD, 's10_tariffs_taxes.py'), encoding='utf-8') as fh:
            src += fh.read()
        loose = [c for c in config.CHOICES if c not in src and c not in stage.DECISIONS_CARRIED]
        self.assertEqual(loose, [])
        self.assertEqual([c for c in stage.DECISIONS_CARRIED if c not in config.CHOICES], [])

    def test_parent_table_stops(self):
        from s10_tariffs import rates as R
        good = R.load_parent_territories()
        self.assertEqual(set(good), {'SMR', 'AND'})
        path = os.path.join(self.tmp, 'p.csv')
        head = 'territory,parent_schedule,preference_reporter,first_chapter,excluded_prefixes,legal_basis,decided_by,decided_on\n'
        row = 'SMR,EUN,EUN,1,72,basis,David,2026-10-23\n'
        for body in (row + row, row.replace('David', ''), row.replace(',1,', ',0,'), row.replace('EUN,EUN', ',EUN')):
            with open(path, 'w') as fh:
                fh.write(head + body)
            with self.assertRaises(ConfigError, msg=body):
                R.load_parent_territories(path)

    def test_read_table_stops(self):
        path = os.path.join(self.tmp, 't.csv')
        with self.assertRaises(ConfigError):
            config.read_table('t.csv', ('a',), path=path)                 # missing
        with open(path, 'w') as fh:
            fh.write('a,b\n1,2\n')
        self.assertEqual(config.read_table('t.csv', ('a',), path=path), [{'a': '1', 'b': '2'}])
        with self.assertRaises(ConfigError):
            config.read_table('t.csv', ('c',), path=path)                 # mis-keyed column
        with self.assertRaises(ConfigError):
            config.read_table('t.csv', ('a',), min_rows=2, path=path)     # too few rows
        with open(path, 'w') as fh:
            fh.write('a,b\n')
        with self.assertRaises(ConfigError):
            config.read_table('t.csv', ('a',), path=path)                 # empty


def fake_build():
    """A two-importer world: AAA (one schedule line) and KOR (one TRQ line)."""
    aaa = Schedule('AAA', 'H6', 2022, 'test')
    aaa.set('010101', 10.0, 'ave')
    aaa.ave['010101'] = 10.0
    kor = Schedule('KOR', 'H6', 2022, 'test')
    kor.set('100590', 3.0, 'reported')
    acc = SimpleNamespace(facing={('AAA', '010101'): 100_000.0}, facing_duty={('AAA', '010101'): 10_000.0},
                          above_mfn=[], no_rate_imp={}, imp_total={'AAA': 100_000.0})
    cm = SimpleNamespace(region_of=lambda iso: {'AAA': 'AAA', 'KOR': 'KOR', 'BBB': 'BBB', 'DEU': 'DEU', 'FRA': 'FRA'}.get(iso, ''))
    rate = {('AAA', 'BBB', 'crp'): 10.0, ('KOR', 'BBB', 'crp'): 3.0}
    duty = {('AAA', 'BBB', 'crp'): (10.0, 100.0), ('KOR', 'BBB', 'crp'): (3.0, 100.0)}
    return SimpleNamespace(acc=acc, schedules={'AAA': aaa, 'KOR': kor}, cm=cm,
                           unions=SimpleNamespace(eu_territory={'DEU', 'FRA'}, fill={}, screen_only={}),
                           wedge_rows=[(100.0, 0.1, 110.0, 10.0)], duty_cell=duty, rate=rate, group_problems=[],
                           C={'trq_fill_measure': config.CHOICES['trq_fill_measure'][0]},
                           state={'log': defaultdict(list), 'cover': {'AAA': {'BBB': [{'partner': 'X'}]}}})


class Screens(unittest.TestCase):
    """Each screen must report its mutation; the hard-check screens must fail on it."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.expl = os.path.join(self.tmp, 'expl.csv')
        self.fta = os.path.join(self.tmp, 'fta.csv')
        self.zip = os.path.join(self.tmp, 'mfn.zip')
        self.write_expl([])
        with open(self.fta, 'w') as fh:
            fh.write('agreement,rta_id,type,importer,exporter,in_force_all_2022,entry_into_force\n')
        zipfile.ZipFile(self.zip, 'w').close()
        self._saved = (suite.EXPLANATIONS, suite.FTA_PAIRS)
        suite.EXPLANATIONS, suite.FTA_PAIRS = self.expl, self.fta

    def tearDown(self):
        suite.EXPLANATIONS, suite.FTA_PAIRS = self._saved

    def write_expl(self, rows):
        with open(self.expl, 'w', newline='') as fh:
            w = csv.writer(fh)
            w.writerow(['importer', 'hs6/sector', 'screen', 'evidence class', 'source file and page', 'decided_by', 'date'])
            w.writerows(rows)

    def run_suite(self, tb, customs=None, prev=None):
        ctx = {'build': tb, 'duty_region': {'AAA': 10.0, 'KOR': 3.0}, 'customs_rows': customs or [],
               'census_usa_musd': 0.0, 'prev_duty': prev, 'mfn_ave_zip': self.zip}
        out = os.path.join(self.tmp, 'out')
        res = suite.run(ctx, report_dir=out, log=lambda *_a: None)
        return res, out

    def rows(self, out, n, key):
        with open(os.path.join(out, f'screen_{n:02d}_{key}.csv'), newline='') as fh:
            return list(csv.DictReader(fh))

    def test_baseline_clean(self):
        res, _out = self.run_suite(fake_build(), prev={('AAA', 'crp'): 10.0, ('KOR', 'crp'): 3.0})
        self.assertEqual(res['material_unexplained'], 0, res['detail'])

    def test_high_rate(self):
        tb = fake_build()
        tb.schedules['AAA'].set('010101', 150.0, 'ave')
        res, _ = self.run_suite(tb)
        self.assertGreater(res['material_unexplained'], 0)
        self.write_expl([['AAA', '010101', 'high_rate', 'statute_confirmed', 'book p.1', 'David', '2026-10-23']])
        res, _ = self.run_suite(tb)
        self.assertEqual(res['material_unexplained'], 0, res['detail'])

    def test_wits_rate_as_given(self):
        tb = fake_build()
        tb.schedules['AAA'].set('010101', 150.0, 'ave')
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w') as inner:
            inner.writestr('a.csv', 'ProductCode,SimpleAverage\n010101,150.00\n')
        with zipfile.ZipFile(self.zip, 'w') as outer:
            outer.writestr('AVEMFN_H6_AAA_2022_U2.zip', buf.getvalue())
        res, out = self.run_suite(tb)
        self.assertEqual(self.rows(out, 2, 'high_rate')[0]['evidence class'], 'wits_rate_as_given')
        self.assertEqual(res['material_unexplained'], 0, res['detail'])
        tb.schedules['AAA'].set('010101', 225.0, 'ave')                   # booked 1.5 x WITS
        res, out = self.run_suite(tb)
        self.assertEqual(self.rows(out, 2, 'high_rate')[0]['evidence class'], 'unexplained')
        self.assertGreater(res['material_unexplained'], 0)

    def test_trains_file_value_and_zambia(self):
        tb = fake_build()
        tb.schedules['AAA'].set('010101', 150.0, 'reported')
        tb.schedules['AAA'].reported['010101'] = 150.0                  # no bulk file: the TRAINS value
        res, out = self.run_suite(tb)
        self.assertEqual(self.rows(out, 2, 'high_rate')[0]['evidence class'], 'wits_rate_as_given')
        tb.schedules['AAA'].set('010101', 225.0, 'reported')            # booked differs from the file
        res, out = self.run_suite(tb)
        self.assertEqual(self.rows(out, 2, 'high_rate')[0]['evidence class'], 'unexplained')
        from s10_tariffs.overrides import ZMB_ADVALOREM
        tb = fake_build()
        zmb = Schedule('ZMB', 'H6', 2022, 'test')
        zmb.set('870421', ZMB_ADVALOREM['870421'], 'class_a_ad_valorem')
        zmb.ave['870421'] = 2621.0
        tb.schedules['ZMB'] = zmb
        tb.acc.facing[('ZMB', '870421')] = 50_000.0
        tb.acc.facing_duty[('ZMB', '870421')] = 25_000.0
        res, out = self.run_suite(tb)
        self.assertEqual(self.rows(out, 8, 'ceiling')[0]['evidence class'], 'tariff_book')
        zmb.set('870421', 40.0, 'class_a_ad_valorem')                   # not the decided rate
        res, out = self.run_suite(tb)
        self.assertEqual(self.rows(out, 8, 'ceiling')[0]['evidence class'], 'unexplained')

    def test_ceiling(self):
        tb = fake_build()
        tb.schedules['AAA'].ave['010101'] = 3000.0
        res, out = self.run_suite(tb)
        self.assertEqual(len(self.rows(out, 8, 'ceiling')), 1)
        self.assertGreater(res['material_unexplained'], 0)

    def test_year_jump(self):
        tb = fake_build()
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w') as inner:
            inner.writestr('a.csv', 'ProductCode,SimpleAverage\n010101,1.5\n')
        with zipfile.ZipFile(self.zip, 'w') as outer:
            outer.writestr('AVEMFN_H6_AAA_2020_U2.zip', buf.getvalue())
        res, out = self.run_suite(tb)
        hit = self.rows(out, 3, 'year_jump')
        self.assertEqual(len(hit), 1)
        self.assertEqual(hit[0]['evidence class'], 'wits_rate_as_given')   # booked equals the TRAINS file
        tb.schedules['AAA'].ave['010101'] = 8.0                             # booked no longer the file value
        res, out = self.run_suite(tb)
        self.assertEqual(self.rows(out, 3, 'year_jump')[0]['evidence class'], 'unexplained')
        self.assertGreater(res['material_unexplained'], 0)

    def test_trq(self):
        tb = fake_build()
        from s10_tariffs.overrides import korea_trq_rates
        decided = korea_trq_rates()['100590']['rate']
        tb.schedules['KOR'].set('100590', decided, 'korea_trq')
        tb.acc.facing[('KOR', '100590')] = 50_000.0
        tb.acc.facing_duty[('KOR', '100590')] = 50_000.0
        res, out = self.run_suite(tb)
        self.assertEqual(self.rows(out, 10, 'trq')[0]['evidence class'], 'trq')   # equals the decided rate
        self.assertEqual(res['material_unexplained'], 0, res['detail'])
        tb.schedules['KOR'].set('100590', decided + 300.0, 'korea_trq')
        tb.acc.facing[('KOR', '100590')] = 50_000.0
        tb.acc.facing_duty[('KOR', '100590')] = 50_000.0
        res, out = self.run_suite(tb)
        self.assertEqual(len(self.rows(out, 10, 'trq')), 1)
        self.assertGreater(res['material_unexplained'], 0)

    def test_identities(self):
        tb = fake_build()
        tb.acc.above_mfn = [('AAA', 'BBB', '010101', 12.0, 10.0)]
        res, _ = self.run_suite(tb)
        self.assertGreater(res['material_unexplained'], 0)
        tb = fake_build()
        tb.wedge_rows = [(100.0, 0.1, 110.0, 11.0)]
        res, _ = self.run_suite(tb)
        self.assertGreater(res['material_unexplained'], 0)

    def test_groups(self):
        tb = fake_build()
        tb.group_problems = ['A36 differs from WITS']
        res, _ = self.run_suite(tb)
        self.assertGreater(res['material_unexplained'], 0)
        tb = fake_build()
        tb.state['log']['unknown_group'].append(('AAA', 'X99', 2021))
        res, _ = self.run_suite(tb)
        self.assertGreater(res['material_unexplained'], 0)

    def test_collections(self):
        customs = [['AAA', '1000', 'src', '', '', 'AAA', 'cash collection, report only (E3)']]
        res, out = self.run_suite(fake_build(), customs=customs)
        hit = self.rows(out, 1, 'collections')[0]
        self.assertEqual(hit['evidence class'], 'gtap_rule')       # no driver above the bar
        self.assertEqual(res['material_unexplained'], 0, res['detail'])
        tb = fake_build()
        tb.schedules['AAA'].set('010101', 150.0, 'ave')               # an unexplained driver above 0.02 bn
        tb.acc.facing_duty[('AAA', '010101')] = 30_000.0
        res, out = self.run_suite(tb, customs=customs)
        self.assertEqual(self.rows(out, 1, 'collections')[0]['evidence class'], 'unexplained')

    def test_diagnostics(self):
        tb = fake_build()
        tb.schedules['AAA'].year = 2015                                # 4 stale
        tb.acc.no_rate_imp = {('AAA', 'no_schedule'): 50_000.0}        # 6 coverage
        tb.rate[('DEU', 'FRA', 'crp')] = 0.0
        tb.unions.screen_only = {'X': {'AAA', 'BBB'}}                  # 11 intra-union
        res, out = self.run_suite(tb, prev={('AAA', 'crp'): 500.0})    # 7 moves
        self.assertEqual(len(self.rows(out, 4, 'stale')), 1)
        self.assertEqual(len(self.rows(out, 6, 'coverage')), 1)
        self.assertGreaterEqual(len(self.rows(out, 7, 'moves')), 1)
        self.assertEqual(len(self.rows(out, 11, 'intra_union')), 1)
        self.assertEqual(res['material_unexplained'], 0, res['detail'])   # diagnostics do not fail

    def test_fta_population(self):
        with open(self.fta, 'a') as fh:
            fh.write('Test FTA,1,FTA,KOR,BBB,True,2010-01-01\nCovered FTA,2,FTA,AAA,BBB,True,2010-01-01\n')
        tb = fake_build()
        tb.duty_cell[('KOR', 'BBB', 'crp')] = (300.0, 1000.0)
        res, out = self.run_suite(tb)
        hits = self.rows(out, 12, 'fta_population')
        self.assertEqual([(h['importer'], h['hs6/sector']) for h in hits], [('KOR', 'BBB')])   # AAA<-BBB is covered
        self.assertEqual(hits[0]['evidence class'], 'missing_schedule_in_source')


class ProductTaxPattern(unittest.TestCase):
    """QA round 2 P1 (product_tax_pattern): VAT weights by agent and commodity, GRS layers, Eurostat control."""
    ICIO = {'crp': 'A01', 'fbt': 'C10T12', 'fin': 'K', 'rle': 'L', 'gov': 'O', 'trd': 'G', 'hea': 'Q'}

    def test_schedules_layers_and_control(self):
        from s10_tariffs import product_tax as PT
        vat = {'DEU': (2022, {'*': 0.19, 'A01': 0.07, 'C10T12': 0.07, 'K': 0.0, 'L': 0.0, 'Q': 0.0})}
        rows, used = PT.pattern_rows(['DEU', 'XAA'], {'DEU': 'single', 'XAA': 'composite'}, self.ICIO, vat)
        w = {(r[0], r[1], r[2]): float(r[3]) for r in rows}
        self.assertEqual(used, {'DEU': 'statutory', 'XAA': 'generic_exemptions'})
        self.assertAlmostEqual(w[('DEU', 'p', 'fbt')], 0.07)
        self.assertAlmostEqual(w[('DEU', 'p', 'trd')], 0.19)
        self.assertEqual(w[('DEU', 'g', 'gov')], 0.0)                              # non-market: outside VAT even with a '*' rate
        self.assertEqual(w[('DEU', 'firm', 'fin')], 0.0)                           # firms: indicator off the exempt services
        self.assertEqual(w[('DEU', 'firm', 'fbt')], 1.0)
        self.assertEqual([w[('XAA', 'p', c)] for c in ('crp', 'fin', 'rle', 'gov', 'hea', 'trd')], [1.0, 0, 0, 0, 0, 1.0])
        grs = {('DEU', 'T_5111'): (2022, 7.0), ('DEU', 'T_5121'): (2022, 1.5), ('USA', 'T_5111'): (2022, 0.0),
               ('USA', 'T_5000'): (2022, 4.5), ('USA', 'T_5121'): (2022, 0.6), ('USA', 'T_5123'): (2022, 0.4)}
        lay = {r[0]: r for r in PT.layer_rows(['DEU', 'USA', 'XAA'], {}, {'DEU': 1.0, 'USA': 1.0}, {'DEU': 1000.0, 'USA': 2000.0, 'XAA': 5.0}, grs)}
        self.assertAlmostEqual(float(lay['DEU'][1]), 70.0)
        self.assertAlmostEqual(float(lay['DEU'][3]), 15.0)
        self.assertAlmostEqual(float(lay['USA'][1]), (4.5 - 0.6 - 0.4) / 100 * 2000.0)   # no VAT: general consumption taxes
        self.assertEqual((lay['XAA'][1], lay['XAA'][3]), ('', ''))
        es = {'DE': (2022, {'CPA_TOTAL': 100.0, 'CPA_A01': 10.0, 'CPA_C10-12': 30.0, 'CPA_H49': -5.0, 'CPA_G45': 20.0, 'CPA_G46': 0.0, 'CPA_G47': 0.0})}
        ctl, neg = PT.control_rows({'DEU'}, {'DEU': 'single'}, es)
        sh = {r[1]: float(r[4]) for r in ctl}
        self.assertAlmostEqual(sum(sh.values()), 1.0)
        self.assertAlmostEqual(sh['fbt'], 0.5)
        self.assertEqual(sh['rdt'], 0.0)                                           # a product subsidy is zero in the control
        self.assertEqual(neg, ['DEU rdt -5.0 MEUR'])
        n, ok, _ = K.product_tax(rows, list(lay.values()), ctl, ['DEU', 'XAA'], len(PT.AGENTS) * len(self.ICIO))
        self.assertTrue(ok, _)
        bad = [r[:3] + ['1.5'] + r[4:] if r[:3] == ['DEU', 'p', 'fbt'] else r for r in rows]
        self.assertFalse(K.product_tax(bad, list(lay.values()), ctl, ['DEU', 'XAA'], len(PT.AGENTS) * len(self.ICIO))[1])
        bad = [r[:3] + ['1'] + r[4:] if r[:3] == ['XAA', 'p', 'hea'] else r for r in rows]
        self.assertFalse(K.product_tax(bad, list(lay.values()), ctl, ['DEU', 'XAA'], len(PT.AGENTS) * len(self.ICIO))[1])

class RoundThreeWE2(unittest.TestCase):
    """Round-3 fix wE2: TRQ fill rule outside KOR and PAN, fuel sales taxes, general-government Gulf taxes."""

    def test_trq_regime_and_rate(self):
        from s10_tariffs import trq
        self.assertEqual(trq.regime_of(0.5), 'in_quota')
        self.assertEqual(trq.regime_of(0.95), 'average')
        self.assertEqual(trq.regime_of(1.2), 'out_of_quota')
        self.assertEqual(trq.regime_of(None), 'in_quota_no_fill')
        st = (2950.82, 0.35, 245.5, 24, 122.95, 0)          # 12 cheese quota pairs (in 0.35, over 245.5)
        rate, status = trq.new_rate(st, 12, ['in_quota'] * 12)
        self.assertEqual(status, 'applied')
        self.assertAlmostEqual(rate, (2950.82 - 12 * 245.5) / 12, places=9)
        rate, _s = trq.new_rate(st, 12, ['out_of_quota'] * 12)
        self.assertAlmostEqual(rate, min((2950.82 - 12 * 0.35) / 12, 245.5), places=9)  # clamped to the HS6 maximum
        rate, _s = trq.new_rate((2946.0, 0.0, 245.5, 24, 122.75, 0), 12, ['average'] * 12, bilateral=True)
        self.assertAlmostEqual(rate, 0.0, places=9)                                       # CUSMA in-quota line is free
        self.assertEqual(trq.new_rate((20.0, 10.0, 10.0, 2, 10.0, 0), 1, ['in_quota'])[1], 'no_pair_structure')
        self.assertEqual(trq.new_rate((20.0, 0.0, 20.0, 2, 10.0, 1), 1, ['in_quota'])[1], 'lines_without_ave')
        self.assertEqual(trq.new_rate((20.0, 0.0, 20.0, 2, 10.0, 0), 2, ['in_quota'] * 2)[1], 'more_quota_lines_than_pairs')
        rate, status = trq.new_rate((66.0, 1.0, 65.0, 2, 33.0, 0), 'all', ['out_of_quota'])   # China maize: whole HS6
        self.assertEqual((status, rate), ('applied', 65.0))
        # trq_bilateral_fill: a lower in-quota minimum, or the same maximum and count with a lower sum
        self.assertTrue(trq.is_bilateral((238.0, 0.0, 238.0, 3, 79.33, 0), (249.37, 3.37, 238.0, 3, 83.12, 0)))
        self.assertTrue(trq.is_bilateral((736.0, 0.0, 249.0, 6, 122.67, 0), (748.36, 0.0, 249.0, 6, 124.73, 0)))
        self.assertFalse(trq.is_bilateral((748.36, 0.0, 249.0, 6, 124.73, 0), (748.36, 0.0, 249.0, 6, 124.73, 0)))
        self.assertFalse(trq.is_bilateral((700.0, 0.0, 200.0, 6, 116.7, 0), (748.36, 0.0, 249.0, 6, 124.73, 0)))
        self.assertFalse(trq.is_bilateral((700.0, 0.0, 249.0, 6, 116.7, 0), None))

    def test_trq_coverage_and_fill(self):
        from s10_tariffs import trq
        lines = {'04069011': [('Q1', 100.0, 50.0, 't', 'doc')], '04069021': [('Q1', 100.0, 50.0, 't', 'doc')],
                 '10011000': [('Q2', 10.0, 10.0, 't', 'doc')], '1006': [('Q3', 10.0, 2.0, 't', 'doc'), ('Q4', 10.0, 9.0, 't', 'doc')]}
        cov = trq.coverage(lines, {'040690', '100111', '100119', '100610', '100620'})
        self.assertEqual(cov['040690'][0], 2)
        self.assertEqual(cov['040690'][1], ['in_quota', 'in_quota'])
        self.assertEqual(cov['100111'][0], 'all')                     # HS2017 1001.10 -> HS2022 1001.11
        self.assertEqual(cov['100111'][1], ['out_of_quota'])
        self.assertEqual(cov['100620'][1], ['in_quota'])               # two quotas on one line: 11 of 20
        self.assertIsNone(trq.line_fill([('Q', 10.0, None, 't', 'd')]))

    def test_trq_check_mutation(self):
        from s10_tariffs import trq
        sch = Schedule('CAN', 'H6', 2022, 'trains_mfn')
        sch.set('040690', 0.4, 'trq_fill')
        rows = [['CAN', '000', 'H6', '040690', 'CANQ013', '12', '24', 'in_quota', '122.95', '0.400000', '0.35', '245.5', 'applied', 's']]
        self.assertTrue(passes(trq.check(rows, {'CAN': sch}, {}, [('CAN', '040690', 0.4)])))
        sch.set('040690', 122.95, 'ave')
        self.assertFalse(passes(trq.check(rows, {'CAN': sch}, {}, [('CAN', '040690', 0.4)])))
        bad = [r[:] for r in rows]
        bad[0][9] = '300.0'
        sch.set('040690', 0.4, 'trq_fill')
        self.assertFalse(passes(trq.check(bad, {'CAN': sch}, {}, [('CAN', '040690', 0.4)])))

    def test_fuel_sales_taxes(self):
        from s10_tariffs import energy_tax
        neer = {(2021, 'IND', 'ROAD', 'GASO'): {'TAXBTJ': 300.0}, (2021, 'IND', 'ROAD', 'DIES'): {'TAXBTJ': 500.0},
                (2021, 'IND', 'INDUSTRY', 'DIES'): {'TAXBTJ': 200.0}, (2021, 'IND', 'ROAD', 'LPG'): {'TAXBTJ': 999.0}}
        tmp = tempfile.mkdtemp()
        try:
            p = os.path.join(tmp, 'fuel_sales_taxes_2022.csv')
            cols = ['region', 'commodity', 'booked_class', 'base_sources', 'value_lcu', 'currency', 'period', 'source', 'url',
                    'retrieved', 'basis', 'decided_by', 'decided_at']
            with open(p, 'w', newline='') as fh:
                w = csv.writer(fh)
                w.writerow(cols)
                w.writerow(['IND', 'pet', 'ROAD', 'GASO;DIES', '80000000000', 'INR', '2022', 'PPAC', 'https://x', 'd', 'b', 'wE2', '2026-10-24'])
            out = energy_tax.fuel_sales_taxes(neer, 'neer_2021', {'IND': 80.0}, path=p)
            self.assertAlmostEqual(out['IND']['total_musd'], 1000.0)
            self.assertAlmostEqual(out['IND']['share'], 0.8)
            self.assertAlmostEqual(out['IND']['booked_musd'], 800.0)
            with self.assertRaises(ConfigError):
                energy_tax.fuel_sales_taxes(neer, 'neer_2021', {}, path=p)          # no exchange rate stops the run
        finally:
            shutil.rmtree(tmp)

    def test_gulf_general_government(self):
        from s10_tariffs import taxes
        S = SimpleNamespace(log=lambda *a: None, flag=lambda *a, **k: None)
        C = {'income_tax_source': 'oecd_grs_2022_then_imf_world_s13_then_imf_gfs_soo_s13_then_world_central',
             'gulf_zakat': 'sau_gfs_g116_other_taxes_incl_zakat_as_cit', 'income_tax_aggregate': 'x', 'tax_missing': 'x',
             'bhr_kwt': 'x', 'world_sector': 'x'}
        world = {('SAU', 'G1112_POGDP_PT_R', 'S1311MIXED'): (2022, 0.53), ('SAU', 'G121_POGDP_PT_R', 'S1311MIXED'): (2022, 0.0),
                 ('SAU', 'G1111_POGDP_PT_R', 'S1311MIXED'): (2022, 0.0),
                 ('ARE', 'G121_POGDP_PT_R', 'S13'): (2022, 2.02), ('ARE', 'G1112_POGDP_PT_R', 'S13'): (2022, 14.6),
                 ('ARE', 'G1111_POGDP_PT_R', 'S13'): (2022, 0.0)}
        gfs = {('SAU', 'G121'): (2022, 2.5), ('SAU', 'G1112'): (2022, 0.53), ('SAU', 'G116'): (2022, 0.6), ('SAU', 'G1111'): (2022, 0.0),
               ('SAU', 'G112'): (2022, 0.0), ('SAU', 'G113'): (2022, 0.01), ('ARE', 'G121'): (2022, 1.48)}
        T = taxes.Taxes(S, C, {}, world, ['SAU', 'ARE'], gfs=gfs)
        T.country_rates()
        self.assertAlmostEqual(T.tax[('SAU', 'ssc')][0], 2.5)
        self.assertAlmostEqual(T.tax[('SAU', 'cit')][0], 1.13)                     # CIT plus zakat
        self.assertAlmostEqual(T.tax[('SAU', 'payroll_property')][0], 0.01)
        self.assertAlmostEqual(T.tax[('ARE', 'ssc')][0], 2.02)                     # WoRLD S13 stays first
        C['income_tax_source'] = 'oecd_grs_2022_then_imf_world_2022'
        T = taxes.Taxes(S, C, {}, world, ['SAU', 'ARE'], gfs=gfs)
        T.country_rates()
        self.assertAlmostEqual(T.tax[('SAU', 'ssc')][0], 0.0)                      # previous rule: central government


if __name__ == '__main__':
    unittest.main()
